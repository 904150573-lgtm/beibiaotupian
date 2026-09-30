"""Extract this label workbook schema and build a self-contained Photoshop job."""
import argparse
import json
import re
import sys
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIELDS = ['产品名称','配料','产品类别','生产日期','保质期','贮存方法','食用方法','温馨提示','致敏物提示','净含量']
NUTRIENTS = ['能量','蛋白质','脂肪','—饱和脂肪','碳水化合物','—糖','钠']

def numeric(ws, cell, visiting=None):
    visiting = set() if visiting is None else visiting
    if cell in visiting:
        raise ValueError(f"{ws.title}!{cell}: 公式循环引用")
    value = ws[cell].value
    if isinstance(value, str) and value.startswith('='):
        # Recompute only the simple source-cell division formulas in this schema.
        m = re.fullmatch(r'=([A-Z]+[1-9][0-9]*)/([0-9]+(?:\.[0-9]+)?)', value)
        if not m:
            raise ValueError(f'{ws.title}!{cell}: 不支持的公式 {value}')
        divisor = Decimal(m[2])
        if not divisor:
            raise ValueError(f"{ws.title}!{cell}: 除数为零")
        source = numeric(ws, m[1], visiting | {cell})
        if source is None:
            raise ValueError(f"{ws.title}!{cell}: 公式引用空值")
        return source / divisor
    if value is None:
        return None
    try:
        number = Decimal(str(value))
    except Exception as error:
        raise ValueError(f"{ws.title}!{cell}: 非数值 {value!r}") from error
    if not number.is_finite() or number < 0:
        raise ValueError(f"{ws.title}!{cell}: 必须为非负有限数值")
    return number

def number_string(value, fmt="General"):
    if fmt != "General":
        m = re.fullmatch(r"0(?:\.(0+))?", fmt)
        if not m:
            raise ValueError(f"不支持的数值格式 {fmt}，请使用 General、0 或 0.00 等")
        dp = len(m[1] or "")
        return format(value.quantize(Decimal(10)**-dp, rounding=ROUND_HALF_UP), f".{dp}f")
    return format(value, 'f').rstrip('0').rstrip('.') if '.' in format(value,'f') else str(value)

def extract(path, schema=None):
    schema = schema or {"columns": {"key":1,"value":3,"unit":4,"nrv":6},"ignored_keys":["产品标签信息","营养成分表"],"basis_values":["每100克","每100g"]}
    cols = schema["columns"]
    wb = load_workbook(path, data_only=False)
    products = []
    try:
        for ws in wb:
            if ws.sheet_state != 'visible':
                continue
            fields, nutrients = {}, {}
            basis = None
            for row in ws:
                key = ws.cell(row[0].row, cols["key"]).value
                if key is None:
                    continue
                key = str(key).strip()
                r = row[0].row
                if key in FIELDS:
                    if key in fields:
                        raise ValueError(f'{ws.title}: 重复字段 {key}')
                    raw = ws.cell(r,cols["value"]).value
                    if raw is not None and str(raw).startswith('='):
                        raise ValueError(f'{ws.title}: 文本字段公式需先转为值 {key}')
                    if ws.cell(r,cols["value"]).data_type == 'e':
                        raise ValueError(f'{ws.title}: Excel 错误值 {key}')
                    if isinstance(raw, (datetime, date)):
                        fmt = ws.cell(r,cols["value"]).number_format
                        if fmt not in ['yyyy-mm-dd','yyyy/mm/dd','yyyy年mm月dd日']:
                            raise ValueError(f'{ws.title}: 日期请使用 yyyy-mm-dd 或文本 {key}')
                        raw = raw.strftime({'yyyy-mm-dd':'%Y-%m-%d','yyyy/mm/dd':'%Y/%m/%d','yyyy年mm月dd日':'%Y年%m月%d日'}[fmt])
                    elif isinstance(raw, (int, float)):
                        raw = number_string(Decimal(str(raw)), ws.cell(r,cols["value"]).number_format)
                    fields[key] = '' if raw is None else str(raw)
                elif key == '项目':
                    if basis is not None:
                        raise ValueError(f'{ws.title}: 重复计量基准')
                    basis = str(ws.cell(r,cols["value"]).value or '')
                elif key in NUTRIENTS:
                    if key in nutrients:
                        raise ValueError(f'{ws.title}: 重复营养项目 {key}')
                    amount = numeric(ws, ws.cell(r,cols["value"]).coordinate)
                    unit = ws.cell(r,cols["unit"]).value
                    expected_unit = '千焦' if key == '能量' else ('毫克' if key == '钠' else '克')
                    if amount is None or unit != expected_unit:
                        raise ValueError(f'{ws.title}: 营养数值或单位缺失 {key}')
                    pct = numeric(ws, ws.cell(r,cols["nrv"]).coordinate)
                    fmt = ws.cell(r,cols["nrv"]).number_format
                    if pct is None:
                        if key != '—糖':
                            raise ValueError(f'{ws.title}: NRV 缺失 {key}')
                        shown = ''
                    else:
                        m = re.fullmatch(r'0(?:\.(0+))?%',fmt)
                        if not m:
                            raise ValueError(f'{ws.title}: 无法准确解释百分比格式 {fmt}')
                        dp = len(m[1] or '')
                        shown = format((pct*100).quantize(Decimal(10)**-dp,rounding=ROUND_HALF_UP), f'.{dp}f')+'%'
                    nutrients[key] = [number_string(amount, ws.cell(r,cols["value"]).number_format)+unit, shown]
                elif key not in schema["ignored_keys"]:
                    raise ValueError(f'{ws.title}: 未映射字段 {key}，先检查，不能静默丢弃')
            for key in FIELDS:
                if key not in ['产品类别','温馨提示'] and not fields.get(key, '').strip():
                    raise ValueError(f'{ws.title}: 必填字段缺失 {key}')
            if set(nutrients) != set(NUTRIENTS) or basis not in schema["basis_values"]:
                raise ValueError(f'{ws.title}: 营养项目或计量基准不同，需重新映射')
            products.append(dict(name=fields['产品名称'],sheet=ws.title,fields=fields,basis=basis,
                amounts='\r'.join(nutrients[n][0] for n in NUTRIENTS),
                percentages='\r'.join(nutrients[n][1] for n in NUTRIENTS),
                warnings=[] if fields.get('产品类别', '').strip() else ['Excel 未提供产品类别，已隐藏对应行，未推断类别。']))
    finally:
        wb.close()
    if not products:
        raise ValueError('没有可处理的产品工作表')
    return products

def build(excel, output, config_path=ROOT/'config/config.json'):
    output, config_path = Path(output), Path(config_path).resolve()
    config = json.loads(config_path.read_text(encoding='utf-8'))
    products = extract(excel, config['excel'])
    mapping = json.loads((config_path.parent/config['template']['mapping']).read_text(encoding='utf-8'))
    ids = {e['id'] for e in mapping}
    if len(ids) != len(mapping):
        raise ValueError('图层映射中 ID 重复')
    rows=config['layout']['rows']
    if sorted(row['field'] for row in rows) != sorted(FIELDS):
        raise ValueError('行映射必须包含每个产品字段且不能重复')
    direct_ids=[i for row in rows if not row.get('duplicate') for i in (row['label_id'],row['value_id'])]
    if len(direct_ids) != len(set(direct_ids)):
        raise ValueError('非复制字段不能共用图层')
    for key in ['names_id','amounts_id','nrv_id','basis_id']:
        if config['layout']['nutrition'][key] not in ids:
            raise ValueError(f'营养表图层不存在：{key}')
    for row in rows:
        if row['field'] not in FIELDS or row['label_id'] not in ids or row['value_id'] not in ids:
            raise ValueError(f'无效行映射 {row}')
    if not isinstance(config['output']['jpeg_quality'], int) or not 0 <= config['output']['jpeg_quality'] <= 12:
        raise ValueError('jpeg_quality 必须是 0–12 的整数')
    job=dict(products=products,template=mapping,config=config,
             width=config['template']['width'],height=config['template']['height'],
             templateByteSize=config['template']['byte_size'])
    engine=(ROOT/'scripts/photoshop_engine.jsx').read_text(encoding='utf-8')
    marker = '/*__INLINE_JOB__*/null'
    if engine.count(marker) != 1 or engine.count('/*__TYPE_HELPERS__*/') != 1:
        raise ValueError('脚本数据插入位置无效')
    generated = engine.replace(marker, json.dumps(job,ensure_ascii=True,indent=2))
    generated = generated.replace('/*__TYPE_HELPERS__*/', (ROOT/'scripts/preserve_type.jsx').read_text(encoding='utf-8'))
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text('#target photoshop\n'+generated,encoding='utf-8')
    manifest=output.with_suffix('.job.json')
    manifest.write_text(json.dumps(job,ensure_ascii=False,indent=2),encoding='utf-8')
    return dict(products=len(products),output=str(output),manifest=str(manifest),warnings=[p['warnings'] for p in products])


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('excel',type=Path)
    ap.add_argument('output',type=Path)
    ap.add_argument('--config',type=Path,default=ROOT/'config/config.json')
    args=ap.parse_args()
    try:
        print(json.dumps(build(args.excel,args.output,args.config),ensure_ascii=False))
    except (ValueError, OSError, KeyError) as error:
        print(f'已停止：{error}',file=sys.stderr)
        return 1
    return 0

if __name__=='__main__':
    sys.exit(main())
