"""Generate a synthetic two-flavour workbook. No real product claims."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def create(path):
    wb = Workbook()
    wb.remove(wb.active)
    for flavour in ['原味', '香辣味']:
        ws = wb.create_sheet(flavour)
        rows = [
            ('产品标签信息', ''), ('产品名称', f'示例脆脆虾（{flavour}）'),
            ('配料', '虾、植物油、食用盐' + ('、辣椒粉' if flavour == '香辣味' else '')),
            ('产品类别', '示例类别'), ('生产日期', '见包装'), ('保质期', '6个月'),
            ('贮存方法', '置于阴凉干燥处，避免阳光直射'), ('食用方法', '开袋即食'),
            ('温馨提示', '开封后请尽快食用'), ('致敏物提示', '本产品含虾及其制品'),
            ('净含量', '50g'), ('营养成分表', ''), ('项目', '每100克')]
        for key, value in rows:
            ws.append([key, None, value])
        for key, value, unit, pct in [
            ('能量', 2000, '千焦', .24), ('蛋白质', 20, '克', .33),
            ('脂肪', 25, '克', .42), ('—饱和脂肪', 5, '克', .25),
            ('碳水化合物', 40, '克', .13), ('—糖', 2.0, '克', None),
            ('钠', 600, '毫克', .30)]:
            ws.append([key, None, value, unit, None, pct])
            ws.cell(ws.max_row, 3).number_format = '0.0' if unit == '克' else '0'
            ws.cell(ws.max_row, 6).number_format = '0%'
        for column in ['A', 'C']:
            ws.column_dimensions[column].width = 24 if column == 'A' else 55
        for column in ['B', 'E']:
            ws.column_dimensions[column].width = 3
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['F'].width = 12
        ws.freeze_panes = 'C2'
        for row in ws:
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
            ws.row_dimensions[row[0].row].height = 30
            row[0].fill = PatternFill('solid', fgColor='EEF2EE')
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == '__main__':
    print(create(Path(__file__).parent / 'products.xlsx'))
