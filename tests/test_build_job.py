import json
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'examples'))
from build_job import build, extract, numeric
from make_example import create


class JobTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = create(Path(self.tmp.name) / 'products.xlsx')

    def change(self, cell, value):
        wb = load_workbook(self.path)
        wb['原味'][cell] = value
        wb.save(self.path)
        wb.close()

    def test_all_flavours_precision_blank_sugar(self):
        products = extract(self.path)
        self.assertEqual([p['sheet'] for p in products], ['原味', '香辣味'])
        self.assertIn('20.0克', products[0]['amounts'])
        self.assertEqual(products[0]['percentages'].split('\r')[5], '')

    def test_unknown_field_rejected(self):
        self.change('A21', '未映射信息')
        with self.assertRaisesRegex(ValueError, '未映射字段'):
            extract(self.path)

    def test_missing_category_hidden_warning(self):
        self.change('C4', None)
        self.assertIn('已隐藏', extract(self.path)[0]['warnings'][0])

    def test_percent_round_half_up(self):
        self.change('F14', .265)
        self.assertEqual(extract(self.path)[0]['percentages'].split('\r')[0], '27%')

    def test_formula_division_and_cycle(self):
        wb = load_workbook(self.path)
        ws = wb['原味']
        ws['F14'] = '=C14/8400'
        self.assertEqual(numeric(ws, 'F14'), numeric(ws, 'C14') / 8400)
        ws['F14'] = '=F14/2'
        with self.assertRaisesRegex(ValueError, '循环'):
            numeric(ws, 'F14')
        wb.close()

    def test_unsupported_formula_and_wrong_unit(self):
        self.change('F14', '=SUM(C14:C15)')
        with self.assertRaisesRegex(ValueError, '不支持的公式'):
            extract(self.path)
        self.change('F14', .24)
        self.change('D14', '克')
        with self.assertRaisesRegex(ValueError, '单位'):
            extract(self.path)

    def test_self_contained_jsx_manifest_and_unicode(self):
        out = Path(self.tmp.name) / 'job.jsx'
        result = build(self.path, out)
        self.assertEqual(result['products'], 2)
        script = out.read_text()
        self.assertTrue(script.startswith('#target photoshop'))
        self.assertNotIn('/*__INLINE_JOB__*/', script)
        self.assertNotIn('/*__TYPE_HELPERS__*/', script)
        manifest = json.loads(out.with_suffix('.job.json').read_text())
        self.assertEqual(manifest['products'][1]['name'], '示例脆脆虾（香辣味）')


if __name__ == '__main__':
    unittest.main()
